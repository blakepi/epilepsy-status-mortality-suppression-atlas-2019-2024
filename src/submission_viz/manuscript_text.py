from __future__ import annotations

import json
import shutil
from datetime import datetime
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from .io import ROOT, ensure_output_dirs, load_config, load_model_frame, primary_estimate, rel_path, write_json
from .tables import add_dataframe_table, build_supplement_tables, build_table1, build_table2, build_table3, build_table4


def _temporal(config: dict) -> pd.DataFrame:
    path = rel_path(config, "temporal_context")
    if path.exists():
        return pd.read_csv(path)
    return pd.read_csv(rel_path(config, "national_year_trend"))


def _references(config: dict) -> pd.DataFrame:
    path = rel_path(config, "references")
    if path.exists():
        return pd.read_csv(path)
    return pd.DataFrame(columns=["number", "reference"])


def captions() -> dict[str, str]:
    return {
        "figure1": "Figure 1. Suppression-aware modeling necessity and public constraint structure. Panel A shows county-year visibility in the public county-year extract. Panel B summarizes how cell-level information and public aggregate totals define the feasible latent-count space. Panel C shows county-period visibility classes. Suppressed county-year cells are treated as known-positive integer intervals from 1–9, not as zeros.",
        "figure2": "Figure 2. Final Wahab HPC Bayesian primary posterior mortality rate ratios. Points show posterior medians and bars show 95% credible intervals. The primary estimand is nonmetro nonadjacent counties compared with large metropolitan counties, adjusted for SVI quartile, county percentage aged ≥65 years, county percentage male, population offset, year effects, and state-level structure.",
        "figure3": "Figure 3. Suppression-handling comparison with the final constrained Bayesian mortality rate ratio estimate. Scenario rows are sensitivity analyses. The Bayesian constrained latent-count model conditions on public county-year, county-period, state-year, national-year, and national-total constraints.",
        "figure4": "Figure 4. Posterior county-level mortality atlas. Panel A maps posterior mean county-period mortality rates per 100,000 person-years. Panel B maps posterior uncertainty width. County-level posterior quantities are model-derived constrained summaries, not observed counts, recovered suppressed counts, or person-level risks.",
        "figureS1": "Supplementary Figure S1. Final all-parameter convergence verification. The card reports the maximum rank-normalized split R-hat and minimum bulk effective sample size across all 69 retained parameters.",
        "figureS2": "Supplementary Figure S2. Posterior mortality-rate-ratio summary for rurality and SVI contrasts.",
        "figureS3": "Supplementary Figure S3. Posterior probability and uncertainty maps for model-derived county-level quantities.",
        "figureS4": "Supplementary Figure S4. Final latent-count constraint validation.",
    }


def _main_markdown(config: dict) -> str:
    primary = primary_estimate(config)
    final = config["final_wahab_handoff"]
    model = load_model_frame(config)
    refs = _references(config)
    ref_lines = "\n".join(f"{int(r.number)}. {r.reference}" for r in refs.itertuples() if pd.notna(r.reference))
    exact = int((model["q002_count_status"] == "exact").sum())
    suppressed = int((model["q002_count_status"] == "suppressed_1_9").sum())
    zeros = int((model["q002_count_status"] == "zero").sum())
    return f"""# {config['project']['title']}

## Abstract

**Objective:** To estimate rurality-associated multiple-cause mortality mentions of epilepsy or status epilepticus in the United States while respecting CDC WONDER small-cell suppression and public aggregate constraints.

**Methods:** We analyzed public CDC WONDER multiple-cause mortality data for ICD-10 G40/G41 mentions from 2019–2024. The final model frame included {final['model_frame_rows']:,} county-year rows from {final['counties']:,} counties. County-year exact counts, explicit zeros, suppressed 1–9 intervals, county-period constraints, state-year totals, national-year totals, and the known national total of {final['national_total_mcod_g40_g41']:,} deaths defined a constrained latent-count system. A negative-binomial Bayesian mortality model estimated mortality rate ratios for rurality and SVI while adjusting for county age and sex composition, population offset, year effects, and state-level structure.

**Results:** The final Wahab HPC production run completed {final['chains']} independent constrained Bayesian chains. The primary nonmetro nonadjacent versus large metropolitan mortality rate ratio was {primary['posterior_median']:.2f} (95% credible interval [CrI], {primary['credible_interval_lower_95']:.2f}–{primary['credible_interval_upper_95']:.2f}); the posterior probability that the mortality rate ratio exceeded 1 was {primary['posterior_probability_gt_1']:.2f}. All {final['parameter_count']} retained parameters passed rank-normalized split R-hat (maximum, {final['maximum_rhat']:.4f}), bulk effective sample size (ESS; minimum, {final['minimum_bulk_ess']:.1f}), and 5%/95% tail ESS (minimum, {final['minimum_tail_ess']:.1f}) criteria. All {final['saved_parameter_draws']:,} saved latent-count draws passed county-year, county-period, state-year, national-year, and grand-total constraint validation.

**Conclusions:** A constrained Bayesian latent-count model provides model-based evidence of higher epilepsy/status epilepticus mortality mention rates in nonmetro nonadjacent counties while preserving the public WONDER suppression structure. Posterior county summaries are model-derived ecological estimates, not recovered suppressed counts or person-level risks.

## Introduction

County-level mortality analyses can inform geographic health surveillance, but CDC WONDER small-cell suppression complicates rare-outcome modeling. Epilepsy and status epilepticus mortality mentions are particularly affected because many county-year cells are small. Visible-only analyses may understate mortality in sparsely populated counties, whereas arbitrary fixed assignments to suppressed cells can distort rurality-associated estimates.

Prior CDC WONDER and epilepsy mortality studies have described rising mortality and geographic disparities, but public county-level suppression limits direct small-area inference [1–3,12–16]. Rurality, social vulnerability, age structure, and sex composition are also geographically patterned, making suppression-aware adjustment essential [4–7]. This study uses public aggregate constraints to model latent county-year mortality counts without using restricted microdata.

The scenario framework remains useful for transparency and sensitivity analysis. The primary inferential layer in this submission is a constrained Bayesian latent-count model that conditions on exact cells, suppressed intervals, county-period constraints, state-year totals, national-year totals, and the reconciled national total.

## Methods

### Data sources and outcome

The outcome was any mention of ICD-10 G40 epilepsy or G41 status epilepticus on the death certificate in CDC WONDER multiple-cause mortality data, 2019–2024 [1,2]. This is interpreted as epilepsy/status epilepticus-related mortality mention, not definitive etiologic attribution. Underlying-cause G40/G41 extracts were retained as contextual analyses [3].

### Rurality, SVI, and covariates

County rurality was collapsed from Rural-Urban Continuum Codes as large metropolitan (RUCC 1), other metropolitan (RUCC 2–3), nonmetropolitan adjacent (RUCC 4, 6, and 8), and nonmetropolitan nonadjacent (RUCC 5, 7, and 9), with large metropolitan counties as the reference category. Social Vulnerability Index (SVI) quartile was included with Q1 as the reference category [4]. County percentage aged ≥65 years, county percentage male, and population offsets were assembled from public county covariates and standardized for modeling [5–7]. Eleven counties without matched covariates were retained and imputed with state medians or modes when available and national values otherwise. Twenty-four Connecticut county-year population offsets missing from the county-year extract were imputed as county-period person-years divided by six; both imputation types were flagged in the model frame.

### Suppression-aware constraint system

Suppressed county-year cells were treated as known-positive integer intervals from 1 to 9, not as zeros. Exact county-year counts and explicit zero cells were fixed. County-period rows imposed exact, zero, or 1–9 interval constraints on six-year county sums. State-year and national-year aggregate totals were enforced exactly, and the 2019–2024 national total was fixed at {final['national_total_mcod_g40_g41']:,} deaths. Counties without direct covariate matches were retained in the latent-count system so aggregate reconciliation was not broken.

### Bayesian constrained model

Latent county-year counts were modeled with a negative-binomial-2 mean linked to the log population offset, rurality, SVI quartile, standardized percentage aged ≥65 years, standardized percentage male, six centered year effects, and 51 centered state/DC effects. The intercept had a normal prior centered on the crude national log rate with standard deviation 5; nonintercept fixed effects had independent normal(0, 1.5) priors. State and year effects had zero-centered normal priors with separate half-normal(0, 1) scale priors, and the negative-binomial shape parameter had a log-normal prior with log mean log(10) and standard deviation 1.5. Posterior county-level summaries are constrained model-derived quantities, not observed or recovered suppressed counts.

### Computation, convergence, and validation

The custom Metropolis-within-Gibbs sampler began from independently seeded feasible allocations generated by mixed-integer linear programming. Constraint-preserving transfer, interval-exploration, swap, and blocked-refresh moves updated latent counts; random-walk Metropolis proposals updated parameter blocks. Eight production chains used seeds 18291–18298 and each ran 300,000 iterations, discarded 75,000 burn-in iterations, and retained every 50th iteration, yielding 4,500 draws per chain and {final['saved_parameter_draws']:,} draws per parameter. The original primary convergence gate passed. A release-time verification then evaluated all {final['parameter_count']} retained parameters using rank-normalized split R-hat, bulk ESS, 5%/95% tail ESS, and Monte Carlo standard errors. All parameters met R-hat ≤1.01 and bulk/tail ESS ≥400; the maximum R-hat was {final['maximum_rhat']:.6f}, minimum bulk ESS was {final['minimum_bulk_ess']:.1f}, and minimum tail ESS was {final['minimum_tail_ess']:.1f}. Supplementary Table S1 reports all parameter diagnostics. Per-chain parameter draws, chain metadata, compressed validation evidence, and derived posterior outputs are included in the reproducibility release. All {final['saved_parameter_draws']:,} saved latent-count draws passed county-year, county-period, state-year, national-year, and grand-total constraint validation.

### Sensitivity and contextual analyses

Visible-only, fixed-value, residual-allocation, and interval-likelihood analyses were retained as sensitivity comparisons. COVID-19 co-mention analyses and underlying-cause analyses were retained as contextual summaries rather than components of the primary Bayesian outcome. Reporting follows observational-study transparency principles, including STROBE guidance [8,9], and a completed STROBE checklist is included in the submission package.

## Results

### Public data structure and constraints

The model frame contained {len(model):,} county-year rows from {model['county_fips'].nunique():,} counties. The county-year extract included {exact:,} exact cells, {suppressed:,} suppressed 1–9 cells, and {zeros:,} explicit zero cells (Table 1; Figure 1). Constraint validation confirmed that saved latent-count draws preserved county-year bounds, county-period constraints, state-year totals, national-year totals, and the {final['national_total_mcod_g40_g41']:,}-death grand total.

### Final Bayesian rurality and SVI associations

The final primary nonmetro nonadjacent versus large metropolitan mortality rate ratio was {primary['posterior_median']:.2f} (95% CrI, {primary['credible_interval_lower_95']:.2f}–{primary['credible_interval_upper_95']:.2f}), with posterior probability that the mortality rate ratio exceeded 1 of {primary['posterior_probability_gt_1']:.2f} (Table 2; Figure 2). Nonmetro adjacent counties also had elevated posterior mortality (1.27, 95% CrI 1.22–1.33). SVI showed a graded association, with SVI Q4 versus Q1 mortality rate ratio 1.43 (95% CrI 1.37–1.50).

### Model-based mortality rates by rurality

Model-derived posterior mortality rates were higher in nonmetro categories than in large metropolitan counties under both conditional and standardized summaries (Table 3). These rates are posterior summaries per 100,000 person-years and should be interpreted as ecological model estimates.

### Comparison with alternative suppression handling

Alternative suppression approaches produced a wide range of estimates (Table 4; Figure 3). Visible-only models and fixed-value stress tests were sensitive to how suppressed positive cells were handled. The final Bayesian constrained model moved beyond fixed assignments by jointly conditioning on county-year intervals, county-period constraints, and public aggregate totals.

### Posterior county atlas and uncertainty

The posterior atlas maps county-period mortality rates and uncertainty (Figure 4). The maps identify spatial heterogeneity in model-derived mortality mention rates while preserving the distinction between posterior estimates and observed public counts.

### Temporal and contextual analyses

National multiple-cause G40/G41 deaths totaled {final['national_total_mcod_g40_g41']:,} during 2019–2024 and increased from 7,583 in 2019 to 10,796 in 2024. COVID-19 co-mention analyses were retained as contextual sensitivity summaries rather than part of the primary Bayesian outcome and were interpreted using CDC death-certificate coding context [17]. Underlying-cause G40/G41 extracts were also retained as context because equivalent public county-year constraints were not available for a parallel primary model.

## Discussion

The principal finding is model-based evidence of higher epilepsy/status epilepticus mortality mention rates in nonmetro nonadjacent counties compared with large metropolitan counties after adjustment for SVI, age structure, sex composition, population offset, year effects, and state-level structure. The posterior estimate was precise in the final constrained Bayesian model, passed the original primary gate, and met the release-time all-parameter convergence criteria.

Suppression handling matters because public county-level mortality data contain many known-positive but hidden small cells. Treating these cells as zero or assigning a single arbitrary value changes the denominator-adjusted rurality pattern and can make sensitivity results look more certain than the public data support.

The Bayesian constrained model adds information by conditioning on public aggregate totals and interval constraints simultaneously. This provides stronger mortality inference than fixed allocation stress tests while retaining the public-data privacy structure. The tradeoff is greater model dependence, so posterior county summaries should be treated as uncertainty-aware ecological estimates rather than recovered suppressed counts.

Limitations include ecological design, death-certificate mention-based outcome definition, possible coding variation, residual confounding by age and comorbidity structure, model dependence, county-equivalent geography issues, and COVID-era disruption. These analyses do not estimate person-level risk or causal effects.

Public aggregate constraints plus constrained Bayesian modeling can extract more information from public CDC WONDER outputs without restricted microdata. The results support rurality-focused mortality surveillance while preserving careful language around suppression, uncertainty, and ecological inference.

## Data and Code Availability

All analyses use public aggregate CDC WONDER outputs and public county covariates. No person-level or restricted data were used. The archived repository release and DOI are cited in the reference list [11]. Posterior county summaries are constrained model-derived estimates and are not observed or recovered suppressed counts.

## Ethics

This study used public deidentified aggregate data and did not involve human-subjects interaction or identifiable private information.

## Generative AI Declaration

Generative AI assistance was used for coding, manuscript editing, and reproducibility package assembly under human direction and in accordance with journal policy expectations [10]. Scientific interpretation, data provenance, and final responsibility remain with the authors.

## References

{ref_lines}
"""


def _supplement_markdown(config: dict) -> str:
    final = config["final_wahab_handoff"]
    temporal = _temporal(config)
    covid_2020_2022 = temporal.loc[temporal["year"].between(2020, 2022), "covid_lower"].sum()
    covid_2020_2024_lower = temporal.loc[temporal["year"].between(2020, 2024), "covid_lower"].sum()
    covid_2020_2024_upper = temporal.loc[temporal["year"].between(2020, 2024), "covid_upper"].sum()
    return f"""# Supplementary Material

## Supplementary Methods

### Data extracts and constraint construction

The constrained Bayesian workflow modeled county-year latent counts while preserving public CDC WONDER suppression intervals and aggregate totals. Exact cells, zero cells, suppressed 1–9 cells, county-period constraints, state-year totals, national-year totals, and the reconciled 2019–2024 total of {final['national_total_mcod_g40_g41']:,} deaths defined the feasible latent-count space.

### Model specification

The primary model used a negative-binomial-2 mortality model with log population offset, rurality, SVI quartile, standardized percentage aged ≥65 years, standardized percentage male, centered year effects, and centered state/DC effects. Large metropolitan counties and SVI Q1 were reference categories. Prior distributions were normal for fixed and random effects, half-normal for state/year scale parameters, and log-normal for the negative-binomial shape parameter, as specified in the main Methods.

### Sampler and validation

Latent counts were updated only through moves that preserved the hard public constraints. Eight chains used seeds 18291–18298, ran 300,000 iterations, discarded 75,000 burn-in iterations, and retained every 50th iteration. Saved latent-count draws were validated against county-year bounds, county-period constraints, state-year totals, national-year totals, and the grand total. Proposal acceptance fractions for some parameter blocks were outside the nominal 0.20–0.45 tuning target (approximately 0.082–0.085 for beta and 0.65–0.94 for log-scale hyperparameters); these are reported as efficiency limitations rather than convergence failures because all rank-normalized R-hat and bulk/tail ESS criteria passed.

### Convergence gate

The final Wahab HPC production run completed {final['chains']} independent constrained Bayesian chains and reported {final['saved_parameter_draws']:,} saved draws per retained parameter. The original primary gate passed. Release-time diagnostics covered all {final['parameter_count']} parameters: maximum rank-normalized split R-hat {final['maximum_rhat']:.6f}, minimum bulk ESS {final['minimum_bulk_ess']:.1f}, and minimum 5%/95% tail ESS {final['minimum_tail_ess']:.1f}.

### Sensitivity and contextual analyses

Sensitivity analyses included visible-only, fixed-value, residual-allocation, and interval-likelihood comparisons. COVID-19 co-mention and underlying-cause extracts were retained as contextual analyses rather than primary Bayesian outcomes.

## Supplementary Results

### Constraint validation

All {final['saved_parameter_draws']:,} saved latent-count draws passed county-year, county-period, state-year, national-year, and grand-total validation checks. Supplementary Table S2 summarizes {final['latent_validation_records']:,} validation records across eight chains.

### Diagnostics

All-parameter convergence diagnostics are reported in Supplementary Table S1. Per-chain parameter draws, chain configurations/statuses, compressed validation evidence, and derived posterior outputs are archived in the reproducibility release; large latent arrays are represented by their cryptographic hashes and complete validation summaries. Supplementary Figure S1 summarizes the final all-parameter verification.

### Sensitivity analyses

Supplementary Tables S3 and S4 provide scenario and sensitivity results. The scenario comparison clarifies which methods preserve the reconciled total, treat suppressed cells as latent positive counts, and provide posterior uncertainty.

### COVID co-mention context

COVID-19 co-mention analyses were contextual and interpreted using CDC death-certificate coding context [17]. The 2020–2022 lower-bound COVID co-mention total was {covid_2020_2022:,.0f} deaths. The broader 2020–2024 bounded urbanization-year context was {covid_2020_2024_lower:,.0f}–{covid_2020_2024_upper:,.0f} deaths. These are different summaries because they use different calendar windows and lower-bound versus bounded definitions.

### Underlying-cause context

Underlying-cause G40/G41 data were retained as contextual summaries because the staged public extracts did not support an equivalently constrained county-year primary Bayesian model.

## Supplementary Tables

Supplementary tables S1–S8 are included below or as accompanying CSV/XLSX files. The full county posterior summary is supplied as CSV/XLSX. Supplementary Table S5 is a data dictionary for the county posterior file rather than a printed county listing.

## Supplementary Figures

Supplementary Figures S1–S4 provide the convergence gate summary, posterior mortality-rate-ratio summary, posterior probability and uncertainty maps, and final latent-count constraint validation.
"""


def _apply_docx_styles(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10
    for style_name, size in [("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 12)]:
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True


def _add_paragraphs_from_markdown(doc: Document, md: str, stop_before: str | None = None) -> None:
    for raw in md.splitlines():
        line = raw.strip()
        if not line:
            continue
        if stop_before and line.startswith(stop_before):
            break
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:])
            run.bold = True
            run.font.size = Pt(16)
        elif line.startswith("## "):
            doc.add_heading(line[3:], level=1)
        elif line.startswith("### "):
            doc.add_heading(line[4:], level=2)
        elif line.startswith("- "):
            doc.add_paragraph(line[2:], style="List Bullet")
        else:
            doc.add_paragraph(line.replace("**", ""))


def _add_figures(doc: Document, config: dict, figure_dir: Path, figure_keys: list[str]) -> None:
    caps = captions()
    names = {
        "figure1": "figure1_suppression_aware_modeling_necessity.png",
        "figure2": "figure2_primary_posterior_irrs.png",
        "figure3": "figure3_suppression_handling_comparison.png",
        "figure4": "figure4_posterior_county_maps.png",
        "figureS1": "figureS1_convergence_gate_card.png",
        "figureS2": "figureS2_final_posterior_mrr_summary.png",
        "figureS3": "figureS3_probability_uncertainty_maps.png",
        "figureS4": "figureS4_constraint_validation_summary.png",
    }
    for index, key in enumerate(figure_keys):
        path = figure_dir / names[key]
        if not path.exists():
            continue
        if index > 0:
            doc.add_page_break()
        doc.add_picture(str(path), width=Inches(6.4))
        cap = doc.add_paragraph(caps[key])
        for run in cap.runs:
            run.italic = True


def build_main_docx(config: dict, out_path: Path) -> None:
    dirs = ensure_output_dirs(config)
    doc = Document()
    _apply_docx_styles(doc)
    md = _main_markdown(config)
    _add_paragraphs_from_markdown(doc, md, stop_before="## References")
    doc.add_page_break()
    doc.add_heading("Tables", level=1)
    add_dataframe_table(doc, "Table 1. Data structure and public aggregate constraints", build_table1(config))
    add_dataframe_table(doc, "Table 2. Final Bayesian primary posterior estimates", build_table2(config), note="Rank-normalized split R-hat, bulk ESS, and 5%/95% tail ESS are derived from all eight production chains. Supplementary Table S1 reports the complete 69-parameter diagnostics.")
    add_dataframe_table(doc, "Table 3. Model-derived posterior mortality rates by rurality", build_table3(config), note="Rates are per 100,000 person-years and are model-derived posterior summaries.")
    doc.add_page_break()
    add_dataframe_table(doc, "Table 4. Suppression-handling comparison", build_table4(config))
    doc.add_page_break()
    doc.add_heading("Figures", level=1)
    _add_figures(doc, config, dirs["figures"] / "main", ["figure1", "figure2", "figure3", "figure4"])
    doc.add_page_break()
    doc.add_heading("References", level=1)
    refs = _references(config)
    for r in refs.itertuples():
        if pd.notna(r.reference):
            doc.add_paragraph(f"{int(r.number)}. {r.reference}")
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def build_supplement_docx(config: dict, out_path: Path) -> None:
    dirs = ensure_output_dirs(config)
    doc = Document()
    _apply_docx_styles(doc)
    _add_paragraphs_from_markdown(doc, _supplement_markdown(config))
    supp_tables = build_supplement_tables(config)
    doc.add_page_break()
    doc.add_heading("Supplementary Tables", level=1)
    notes = {
        "S1_diagnostics": "All 69 retained parameters are reported. R-hat is rank-normalized and folded; ESS values are multi-chain bulk and 5%/95% tail estimates. Release criteria were R-hat ≤1.01 and bulk/tail ESS ≥400.",
        "S2_constraint_validation": "Production validation contains 489,696 records across eight chains, 12 checks, and 40,808 chain-specific allocation/draw labels; all failed-record counts are zero.",
        "S5_county_posterior_dictionary": "The full county posterior summary is supplied as CSV/XLSX. The DOCX supplement prints the data dictionary only to avoid an unreadable county listing.",
    }
    for title, df in supp_tables.items():
        add_dataframe_table(doc, title, df, note=notes.get(title))
    doc.add_page_break()
    doc.add_heading("Supplementary Figures", level=1)
    _add_figures(doc, config, dirs["figures"] / "supplement", ["figureS1", "figureS2", "figureS3", "figureS4"])
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def write_captions_file(config: dict) -> Path:
    dirs = ensure_output_dirs(config)
    path = dirs["manuscript"] / "figure_captions.md"
    lines = ["# Figure Captions", ""]
    for key, text in captions().items():
        lines.extend([f"## {key}", "", text, ""])
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def write_data_code_availability(config: dict) -> Path:
    dirs = ensure_output_dirs(config)
    path = dirs["manuscript"] / "data_code_availability.md"
    text = """# Data and Code Availability

This analysis uses public aggregate CDC WONDER mortality outputs and public county covariates. No person-level or restricted data are included. Posterior county-level estimates are constrained Bayesian model summaries, not observed or recovered suppressed counts.
"""
    path.write_text(text, encoding="utf-8")
    return path


def write_run_report(config: dict) -> Path:
    primary = primary_estimate(config)
    dirs = ensure_output_dirs(config)
    report = ROOT / "RUN_BAYES_CONSTRAINED_REPORT.md"
    text = f"""# Bayesian Constrained Final Run Report

Generated: {datetime.now().isoformat(timespec='seconds')}

## Final Wahab HPC Result

The final Wahab HPC production run completed eight independent constrained Bayesian chains. All 69 retained parameters passed rank-normalized split R-hat and bulk/tail ESS criteria. The maximum R-hat was {config['final_wahab_handoff']['maximum_rhat']:.6f}; minimum bulk ESS was {config['final_wahab_handoff']['minimum_bulk_ess']:.1f}; minimum 5%/95% tail ESS was {config['final_wahab_handoff']['minimum_tail_ess']:.1f}. All 36,000 saved latent-count draws passed constraint validation.

- National MCOD G40/G41 total: {config['final_wahab_handoff']['national_total_mcod_g40_g41']:,}
- Model frame: {config['final_wahab_handoff']['model_frame_rows']:,} county-year rows from {config['final_wahab_handoff']['counties']:,} counties
- Saved parameter draws in final summaries: {config['final_wahab_handoff']['saved_parameter_draws']:,}
- Primary nonmetro nonadjacent vs large metropolitan mortality rate ratio: {primary['posterior_median']:.2f} (95% CrI {primary['credible_interval_lower_95']:.2f}–{primary['credible_interval_upper_95']:.2f}); Pr(IRR > 1)={primary['posterior_probability_gt_1']:.2f}

## Outputs

- Final submission manuscript: `{(dirs['package'] / 'manuscript_submission_FINAL.docx').relative_to(ROOT)}`
- Final submission supplement: `{(dirs['package'] / 'supplement_submission_FINAL.docx').relative_to(ROOT)}`
- Final QC report: `{(dirs['qc'] / 'SUBMISSION_QC_REPORT_FINAL.md').relative_to(ROOT)}`

## Interpretation Guardrails

The outcome is a multiple-cause mortality mention of ICD-10 G40/G41, not person-level incidence or etiologic attribution. Posterior county summaries are model-derived constrained Bayesian quantities and should not be described as recovered true suppressed counts.
"""
    report.write_text(text, encoding="utf-8")
    return report


def build_main_manuscript(config: dict | None = None) -> dict[str, Path]:
    config = config or load_config()
    dirs = ensure_output_dirs(config)
    md = dirs["manuscript"] / "manuscript_submission.md"
    docx = dirs["manuscript"] / "manuscript_submission.docx"
    md.write_text(_main_markdown(config), encoding="utf-8")
    build_main_docx(config, docx)
    write_captions_file(config)
    write_data_code_availability(config)
    write_run_report(config)
    shutil.copy2(docx, ROOT / "manuscript/manuscript_main_bayes_constrained.docx")
    (ROOT / "manuscript/manuscript_main_bayes_constrained.md").write_text(_main_markdown(config), encoding="utf-8")
    return {"markdown": md, "docx": docx}


def build_supplement(config: dict | None = None) -> dict[str, Path]:
    config = config or load_config()
    dirs = ensure_output_dirs(config)
    md = dirs["manuscript"] / "supplement_submission.md"
    docx = dirs["manuscript"] / "supplement_submission.docx"
    md.write_text(_supplement_markdown(config), encoding="utf-8")
    build_supplement_docx(config, docx)
    shutil.copy2(docx, ROOT / "supplement/supplement_bayes_constrained.docx")
    (ROOT / "supplement/supplement_bayes_constrained.md").write_text(_supplement_markdown(config), encoding="utf-8")
    return {"markdown": md, "docx": docx}


def write_reproducibility_manifest(config: dict) -> Path:
    dirs = ensure_output_dirs(config)
    manifest = {
        "created": datetime.now().isoformat(timespec="seconds"),
        "config": str(Path(config["_config_path"]).relative_to(ROOT)),
        "final_wahab_hpc": {
            "chains": config["final_wahab_handoff"]["chains"],
            "saved_parameter_draws": config["final_wahab_handoff"]["saved_parameter_draws"],
            "primary_rhat": config["final_wahab_handoff"]["primary_rhat"],
            "primary_ess": config["final_wahab_handoff"]["primary_ess"],
            "parameter_count": config["final_wahab_handoff"]["parameter_count"],
            "maximum_rhat": config["final_wahab_handoff"]["maximum_rhat"],
            "minimum_bulk_ess": config["final_wahab_handoff"]["minimum_bulk_ess"],
            "minimum_tail_ess": config["final_wahab_handoff"]["minimum_tail_ess"],
            "primary_mortality_rate_ratio": "1.23 (95% CrI 1.17–1.30)",
        },
        "commands": [
            ".\\.venv\\Scripts\\python.exe scripts\\50_build_submission_figures.py",
            ".\\.venv\\Scripts\\python.exe scripts\\51_build_submission_tables.py",
            ".\\.venv\\Scripts\\python.exe scripts\\52_polish_manuscript_submission.py",
            ".\\.venv\\Scripts\\python.exe scripts\\53_build_submission_supplement.py",
            ".\\.venv\\Scripts\\python.exe scripts\\55_package_submission_assets.py",
            ".\\.venv\\Scripts\\python.exe scripts\\54_qc_submission_package.py",
        ],
    }
    path = dirs["package"] / "reproducibility_manifest.json"
    write_json(path, manifest)
    return path
