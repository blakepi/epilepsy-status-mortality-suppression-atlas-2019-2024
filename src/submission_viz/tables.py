from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.shared import Inches, Pt

from .io import ROOT, ensure_output_dirs, final_estimates, load_config, load_county_summary, load_model_frame, rel_path

EM_DASH = "—"


def _write_csv_xlsx(df: pd.DataFrame, stem: str, out_dir: Path) -> list[Path]:
    csv_path = out_dir / f"{stem}.csv"
    xlsx_path = out_dir / f"{stem}.xlsx"
    df.to_csv(csv_path, index=False)
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="table")
    return [csv_path, xlsx_path]


def _write_table_preview_png(df: pd.DataFrame, stem: str, out_dir: Path, max_rows: int = 16) -> Path:
    import matplotlib.pyplot as plt

    preview = df.head(max_rows).astype(str)
    width = max(8.5, min(16, 1.7 * len(preview.columns)))
    height = max(2.8, min(12, 0.42 * (len(preview) + 2)))
    fig, ax = plt.subplots(figsize=(width, height), dpi=180)
    ax.axis("off")
    table = ax.table(
        cellText=preview.values,
        colLabels=preview.columns,
        loc="center",
        cellLoc="left",
        colLoc="left",
    )
    table.auto_set_font_size(False)
    table.set_fontsize(7.5)
    table.scale(1, 1.22)
    for (row, _col), cell in table.get_celld().items():
        cell.set_edgecolor("#D0D7DE")
        if row == 0:
            cell.set_facecolor("#EAF2F8")
            cell.set_text_props(weight="bold", color="#1B2631")
        else:
            cell.set_facecolor("#FFFFFF" if row % 2 else "#F8FAFC")
    out_path = out_dir / f"{stem}.png"
    fig.savefig(out_path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return out_path


def _fmt_int(value) -> str:
    return f"{int(float(value)):,}"


def _fmt_est(value, digits: int) -> str:
    if pd.isna(value):
        return EM_DASH
    return f"{float(value):.{digits}f}"


def _fmt_interval(low, high, digits: int) -> str:
    return f"{_fmt_est(low, digits)}–{_fmt_est(high, digits)}"


def _fmt_optional_number(value, digits: int = 1) -> str:
    if pd.isna(value):
        return EM_DASH
    return f"{float(value):,.{digits}f}"


def _clean_contrast(term: str) -> str:
    return {
        "metro_other vs metro_large": "Other metropolitan vs large metropolitan",
        "nonmetro_adjacent vs metro_large": "Nonmetro adjacent vs large metropolitan",
        "nonmetro_nonadjacent vs metro_large": "Nonmetro nonadjacent vs large metropolitan",
        "SVI Q2 vs Q1": "SVI Q2 vs Q1",
        "SVI Q3 vs Q1": "SVI Q3 vs Q1",
        "SVI Q4 vs Q1": "SVI Q4 vs Q1",
        "pct_age65 coefficient": "Percentage aged ≥65 years, standardized",
        "pct_male coefficient": "Percentage male, standardized",
    }.get(term, term)


def build_table1(config: dict) -> pd.DataFrame:
    model = load_model_frame(config)
    audit = pd.read_csv(rel_path(config, "data_audit"))
    audit_map = dict(zip(audit["metric"], audit["value"]))
    rows = [
        ["Model frame", "County-year rows", len(model), "One row per county-year, 2019–2024."],
        ["Model frame", "Counties", model["county_fips"].nunique(), "No county was dropped from the constraint system."],
        ["Model frame", "State/DC equivalents", model["state_fips"].nunique(), "State-year constraints were applied across represented jurisdictions."],
        ["County-year visibility", "Exact cells", (model["q002_count_status"] == "exact").sum(), "Observed counts fixed exactly."],
        ["County-year visibility", "Suppressed cells", (model["q002_count_status"] == "suppressed_1_9").sum(), "Latent integer count constrained to 1–9."],
        ["County-year visibility", "Explicit zero cells", (model["q002_count_status"] == "zero").sum(), "Latent count fixed at 0."],
        ["County-period visibility", "Exact rows", audit_map.get("county_period_exact_rows", 0), "County-period sums fixed exactly."],
        ["County-period visibility", "Suppressed rows", audit_map.get("county_period_suppressed_rows", 0), "County-period sums constrained to 1–9."],
        ["County-period visibility", "Zero rows", audit_map.get("county_period_zero_rows", 0), "County-period sums fixed at 0."],
        ["Aggregate constraints", "National total", config["final_wahab_handoff"]["national_total_mcod_g40_g41"], "Known 2019–2024 MCOD G40/G41 total."],
        ["Aggregate constraints", "National-year constraints", model["year"].nunique(), "Each year reconciles to the national-year extract."],
        ["Aggregate constraints", "State-year constraints", model[["state_fips", "year"]].drop_duplicates().shape[0], "Each state-year reconciles to the state-year extract."],
        ["Covariates", "Unmatched covariate counties", int(model.groupby("county_fips")["unmatched_covariate_flag"].max().sum()), "Kept in constraints; imputed for modeling as documented."],
    ]
    df = pd.DataFrame(rows, columns=["Domain", "Quantity", "Value", "Notes"])
    df["Value"] = df["Value"].map(_fmt_int)
    return df


def build_table2(config: dict) -> pd.DataFrame:
    df = final_estimates(config).copy()
    rows = []
    for _, row in df.iterrows():
        is_irr = row["estimate_type"] == "IRR"
        digits = 2 if is_irr else 3
        rows.append(
            {
                "Family": row["family"],
                "Contrast or coefficient": _clean_contrast(row["term"]),
                "Scale": "Mortality rate ratio" if is_irr else "Log mortality-rate coefficient",
                "Posterior median": _fmt_est(row["posterior_median"], digits),
                "95% CrI": _fmt_interval(row["credible_interval_lower_95"], row["credible_interval_upper_95"], digits),
                "Pr(IRR > 1)": _fmt_est(row.get("posterior_probability_gt_1"), 2) if is_irr else EM_DASH,
                "R-hat": _fmt_est(row.get("r_hat"), 3),
                "Bulk ESS": _fmt_est(row.get("ess_bulk"), 1),
                "Tail ESS": _fmt_est(row.get("ess_tail"), 1),
            }
        )
    return pd.DataFrame(rows)


def build_table3(config: dict) -> pd.DataFrame:
    rates = pd.read_csv(rel_path(config, "adjusted_rates_table")).copy()
    rates["Rurality category"] = rates["rurality"].map(
        {
            "metro_large": "Large metropolitan",
            "metro_other": "Other metropolitan",
            "nonmetro_adjacent": "Nonmetro adjacent",
            "nonmetro_nonadjacent": "Nonmetro nonadjacent",
        }
    ).fillna(rates["rurality"])
    rates["Estimate type"] = rates["estimate_type"].str.replace("_", " ", regex=False)
    rates["Posterior median rate"] = rates["posterior_median_rate_per_100k"].map(lambda x: f"{x:.2f}")
    rates["95% CrI"] = rates.apply(
        lambda r: f"{r['credible_interval_lower_95']:.2f}–{r['credible_interval_upper_95']:.2f}",
        axis=1,
    )
    return rates[["Rurality category", "Estimate type", "Posterior median rate", "95% CrI"]]


def _full_scenario_table(config: dict) -> pd.DataFrame:
    prior = pd.read_csv(rel_path(config, "prior_scenario_table")).copy()
    out = prior.rename(
        columns={
            "Scenario": "Method",
            "Preserves reconciled total?": "Preserves reconciled total",
            "95% CI": "95% interval",
        }
    )
    if "Tier" in out.columns:
        out = out.drop(columns=["Tier"])
    out["Models latent counts probabilistically"] = "No"
    out["Provides posterior uncertainty"] = "No"
    out["Treats suppressed cells as latent positive counts"] = "No"
    latent_partial = out["Method"].str.contains("residual allocation|Interval-likelihood", case=False, regex=True)
    out.loc[latent_partial, "Treats suppressed cells as latent positive counts"] = "Partly"
    out.loc[out["Method"].eq("Interval-likelihood negative binomial"), "Models latent counts probabilistically"] = "Partly"
    bayes = {
        "Method": "Bayesian constrained latent-count model",
        "Assigned deaths, full extract": "58,380",
        "Analytic model deaths": "58,380 latent-count system",
        "Preserves reconciled total": "Yes",
        "IRR": "1.23",
        "95% interval": "1.17–1.30",
        "Treats suppressed cells as latent positive counts": "Yes",
        "Models latent counts probabilistically": "Yes",
        "Provides posterior uncertainty": "Yes",
    }
    out = pd.concat([out, pd.DataFrame([bayes])], ignore_index=True)
    out["95% interval"] = out["95% interval"].astype(str).str.replace("-", "–", regex=False)
    return out


def build_table4(config: dict) -> pd.DataFrame:
    full = _full_scenario_table(config)
    return full[
        [
            "Method",
            "Preserves reconciled total",
            "Treats suppressed cells as latent positive counts",
            "Models latent counts probabilistically",
            "Provides posterior uncertainty",
            "IRR",
            "95% interval",
        ]
    ].rename(columns={"IRR": "Nonmetro nonadjacent mortality rate ratio"})


def _constraint_summary(config: dict) -> pd.DataFrame:
    constraints = pd.read_csv(rel_path(config, "constraint_validation"))
    required = {"check", "labels_checked", "validation_records", "failed_records"}
    if not required.issubset(constraints.columns):
        raise ValueError(f"Production constraint summary is missing columns: {sorted(required - set(constraints.columns))}")
    summary = constraints.groupby("check", as_index=False).agg(
        labels_checked=("labels_checked", "sum"),
        validation_records=("validation_records", "sum"),
        failures=("failed_records", "sum"),
    ).sort_values("check")
    labels = {
        "all_counts_integer": "All counts integer",
        "county_period_constraints_respected": "County-period constraints respected",
        "county_year_bounds_respected": "County-year bounds respected",
        "exact_county_year_counts_unchanged": "Exact county-year counts unchanged",
        "explicit_zero_county_year_counts_unchanged": "Explicit-zero county-year counts unchanged",
        "grand_total_equals_58380": "Grand total equals 58,380",
        "grand_total_matches_modeled_years": "Grand total matches modeled years",
        "national_year_totals_equal_q003": "National-year totals equal Q003",
        "no_counties_silently_dropped": "No counties silently dropped",
        "no_negative_counts": "No negative counts",
        "row_count_matches_frame": "Row count matches model frame",
        "state_year_totals_equal_q004": "State-year totals equal Q004",
        "suppressed_county_year_counts_remain_1_9": "Suppressed county-year counts remain 1–9",
    }
    summary["Validation check"] = summary["check"].map(labels).fillna(summary["check"].str.replace("_", " ", regex=False))
    return summary[["Validation check", "labels_checked", "validation_records", "failures"]].rename(
        columns={"labels_checked": "Units checked", "validation_records": "Validation records", "failures": "Failed records"}
    )


def _diagnostics_table(config: dict) -> pd.DataFrame:
    diagnostics = pd.read_csv(rel_path(config, "parameter_diagnostics"))
    required = {"parameter", "label", "r_hat", "ess_bulk", "ess_tail", "mcse_mean", "chains", "draws_per_chain"}
    if not required.issubset(diagnostics.columns):
        raise ValueError(f"Production diagnostics are missing columns: {sorted(required - set(diagnostics.columns))}")
    diagnostics["Parameter family"] = diagnostics["parameter"].map(
        lambda value: (
            "State effect" if value.startswith("state_effect[")
            else "Year effect" if value.startswith("year_effect[")
            else "Rurality" if value.startswith("primary_rurality_")
            else "SVI" if value.startswith("svi_quartile_")
            else "Hyperparameter" if value in {"kappa", "sigma_state", "sigma_year"}
            else "Fixed effect"
        )
    )
    diagnostics["Diagnostic status"] = np.where(
        (diagnostics["r_hat"] <= 1.01)
        & (diagnostics["ess_bulk"] >= 400)
        & (diagnostics["ess_tail"] >= 400),
        "Pass",
        "Fail",
    )
    return pd.DataFrame(
        {
            "Parameter family": diagnostics["Parameter family"],
            "Tracked quantity": diagnostics["label"].map(_clean_contrast),
            "R-hat": diagnostics["r_hat"].map(lambda value: f"{value:.6f}"),
            "Bulk ESS": diagnostics["ess_bulk"].map(lambda value: f"{value:.1f}"),
            "Tail ESS": diagnostics["ess_tail"].map(lambda value: f"{value:.1f}"),
            "MCSE mean": diagnostics["mcse_mean"].map(lambda value: f"{value:.6g}"),
            "Chains": diagnostics["chains"].astype(int),
            "Draws per chain": diagnostics["draws_per_chain"].astype(int),
            "Diagnostic status": diagnostics["Diagnostic status"],
        }
    )


def _format_sensitivity_table(df: pd.DataFrame) -> pd.DataFrame:
    if df.empty:
        return df
    out = df.copy()
    rename = {
        "sensitivity": "Sensitivity",
        "category": "Category",
        "counties": "Counties",
        "posterior_mean_deaths": "Posterior mean deaths",
        "person_years": "Person-years",
        "rate_per_100k": "Rate per 100,000",
        "note": "Note",
    }
    out = out.rename(columns=rename)
    sensitivity_labels = {
        "rurality_only_primary_grouping": "Rurality-only primary grouping",
        "nchs_urban_rural_sensitivity": "NCHS urban-rural sensitivity",
        "binary_rucc_sensitivity": "Binary RUCC sensitivity",
        "covariate_unmatched_excluded_description": "Covariate-unmatched sensitivity",
        "underlying_cause_context": "Underlying-cause context",
        "covid_context": "COVID context",
    }
    category_labels = {
        "metro_large": "Large metropolitan",
        "metro_other": "Other metropolitan",
        "nonmetro_adjacent": "Nonmetro adjacent",
        "nonmetro_nonadjacent": "Nonmetro nonadjacent",
        "large_central_metro": "Large central metropolitan",
        "large_fringe_metro": "Large fringe metropolitan",
        "medium_metro": "Medium metropolitan",
        "small_metro": "Small metropolitan",
        "micropolitan": "Micropolitan",
        "noncore": "Noncore",
        "RUCC_1_3": "RUCC 1–3",
        "RUCC_4_9": "RUCC 4–9",
        "False": "Covariate-unmatched counties",
        "True": "Covariate-matched counties",
        "parallel constrained model not run": "Parallel constrained model not run",
        "2020-2022 COVID co-mention lower-bound deaths": "2020–2022 COVID co-mention lower-bound deaths",
    }
    if "Sensitivity" in out:
        out["Sensitivity"] = out["Sensitivity"].map(lambda x: sensitivity_labels.get(str(x), str(x).replace("_", " ").title()))
    if "Category" in out:
        out["Category"] = out["Category"].map(lambda x: EM_DASH if pd.isna(x) else category_labels.get(str(x), str(x).replace("_", " ").title()))
    if "Counties" in out:
        out["Counties"] = out["Counties"].map(lambda x: _fmt_optional_number(x, 0))
    if "Posterior mean deaths" in out:
        out["Posterior mean deaths"] = out["Posterior mean deaths"].map(lambda x: _fmt_optional_number(x, 1))
    if "Person-years" in out:
        out["Person-years"] = out["Person-years"].map(lambda x: _fmt_optional_number(x, 0))
    if "Rate per 100,000" in out:
        out["Rate per 100,000"] = out["Rate per 100,000"].map(lambda x: _fmt_optional_number(x, 2))
    if "Note" in out:
        out["Note"] = out["Note"].fillna(EM_DASH)
    return out


def _county_dictionary() -> pd.DataFrame:
    rows = [
        ("county_fips", "Five-digit county FIPS code."),
        ("county_name", "County or county-equivalent name."),
        ("state_name", "State name."),
        ("posterior_mean_county_period_deaths", "Posterior mean model-derived county-period deaths."),
        ("posterior_median_county_period_deaths", "Posterior median model-derived county-period deaths."),
        ("deaths_credible_interval_lower_95", "Lower 95% credible limit for county-period deaths."),
        ("deaths_credible_interval_upper_95", "Upper 95% credible limit for county-period deaths."),
        ("posterior_mean_rate_per_100k", "Posterior mean county-period rate per 100,000 person-years."),
        ("rate_credible_interval_lower_95", "Lower 95% credible limit for rate."),
        ("rate_credible_interval_upper_95", "Upper 95% credible limit for rate."),
        ("posterior_probability_rate_exceeds_national_rate", "Posterior probability that county rate exceeds the national rate."),
        ("original_q001_status", "Public county-period visibility status."),
        ("rurality", "Collapsed rurality category."),
        ("svi_quartile", "Social Vulnerability Index quartile."),
    ]
    return pd.DataFrame(rows, columns=["Column", "Description"])


def build_supplement_tables(config: dict) -> dict[str, pd.DataFrame]:
    model = load_model_frame(config)
    county = load_county_summary(config).copy()
    county["county_fips"] = county["county_fips"].astype(str).str.zfill(5)
    temporal = pd.read_csv(rel_path(config, "temporal_context"))
    covid_context = temporal[["year", "deaths", "rate_per_100k", "covid_lower", "covid_upper", "covid_pct_of_mcod"]].copy()
    ucd = pd.read_csv(rel_path(config, "ucd_profile"))
    data_audit = pd.DataFrame(
        [
            ["county_year_rows", len(model)],
            ["counties", model["county_fips"].nunique()],
            ["national_total_constraint", config["final_wahab_handoff"]["national_total_mcod_g40_g41"]],
            ["saved_parameter_draws_final_summary", config["final_wahab_handoff"]["saved_parameter_draws"]],
            ["chains", config["final_wahab_handoff"]["chains"]],
        ],
        columns=["Metric", "Value"],
    )
    sensitivity_path = ROOT / "tables/bayes_constrained_sensitivity_summaries.csv"
    sensitivity = pd.read_csv(sensitivity_path) if sensitivity_path.exists() else pd.DataFrame()
    return {
        "S1_diagnostics": _diagnostics_table(config),
        "S2_constraint_validation": _constraint_summary(config),
        "S3_scenario_comparison_full": _full_scenario_table(config),
        "S4_sensitivity_summaries": _format_sensitivity_table(sensitivity),
        "S5_county_posterior_dictionary": _county_dictionary(),
        "S6_data_audit": data_audit,
        "S7_covid_context": covid_context,
        "S8_underlying_cause_context": ucd,
    }


def _set_table_style(table) -> None:
    table.style = "Table Grid"
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                paragraph.paragraph_format.space_before = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.2 if row_idx else 8.7)
                    run.font.name = "Calibri"
                    if row_idx == 0:
                        run.bold = True


def add_dataframe_table(doc: Document, title: str, df: pd.DataFrame, max_rows: int | None = None, note: str | None = None) -> None:
    doc.add_paragraph(title, style="Heading 2")
    shown = df if max_rows is None else df.head(max_rows)
    table = doc.add_table(rows=1, cols=len(shown.columns))
    for i, col in enumerate(shown.columns):
        table.rows[0].cells[i].text = str(col)
    for _, row in shown.iterrows():
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = "" if pd.isna(value) else str(value)
    _set_table_style(table)
    if max_rows is not None and len(df) > max_rows:
        doc.add_paragraph(f"Showing first {max_rows:,} of {len(df):,} rows; full CSV/XLSX table is included in the package.")
    if note:
        note_paragraph = doc.add_paragraph(note)
        for run in note_paragraph.runs:
            run.italic = True
            run.font.size = Pt(9)


def build_publication_tables_docx(tables: dict[str, pd.DataFrame], supplement: dict[str, pd.DataFrame], out_path: Path) -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading("Submission Tables", level=1)
    for title, df in tables.items():
        add_dataframe_table(doc, title, df)
    doc.add_page_break()
    doc.add_heading("Supplementary Tables", level=1)
    for title, df in supplement.items():
        add_dataframe_table(doc, title, df, max_rows=10 if "preview" in title.lower() else None)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)


def build_all_tables(config: dict | None = None) -> dict[str, Path]:
    config = config or load_config()
    dirs = ensure_output_dirs(config)
    out_dir = dirs["tables"]
    for path in out_dir.glob("*"):
        if path.is_file() and path.suffix.lower() in {".csv", ".xlsx", ".docx", ".png"}:
            path.unlink()
    main = {
        "Table 1. Data structure and public aggregate constraints": build_table1(config),
        "Table 2. Final Bayesian primary posterior estimates": build_table2(config),
        "Table 3. Model-derived posterior mortality rates by rurality": build_table3(config),
        "Table 4. Suppression-handling comparison": build_table4(config),
    }
    stems = [
        "table1_data_structure_constraints",
        "table2_primary_posterior_estimates",
        "table3_adjusted_rates_by_rurality",
        "table4_suppression_handling_comparison",
    ]
    written: list[Path] = []
    for stem, (_, df) in zip(stems, main.items()):
        written.extend(_write_csv_xlsx(df, stem, out_dir))
        written.append(_write_table_preview_png(df, stem, out_dir))
    supp = build_supplement_tables(config)
    for stem, df in supp.items():
        written.extend(_write_csv_xlsx(df, stem.lower(), out_dir))
    workbook = out_dir / "submission_tables_workbook.xlsx"
    with pd.ExcelWriter(workbook, engine="openpyxl") as writer:
        for title, df in {**main, **supp}.items():
            df.to_excel(writer, index=False, sheet_name=title[:31].replace("/", "-"))
    written.append(workbook)
    docx_path = out_dir / "publication_tables.docx"
    build_publication_tables_docx(main, supp, docx_path)
    written.append(docx_path)
    county_full = rel_path(config, "county_posterior_summary")
    if county_full.exists():
        csv_target = out_dir / "supplement_s5_county_posterior_summary_full.csv"
        xlsx_target = out_dir / "supplement_s5_county_posterior_summary_full.xlsx"
        shutil.copy2(county_full, csv_target)
        pd.read_csv(county_full).to_excel(xlsx_target, index=False)
        written.extend([csv_target, xlsx_target])
    return {path.name: path for path in written}
